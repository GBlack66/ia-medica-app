import streamlit as st
from google import genai
import PyPDF2
import io
from pptx import Presentation

# 1. Configuración de la página
st.set_page_config(page_title="Generador de Presentaciones Médicas", layout="wide")
st.title("📊 Generador Pro de Presentaciones (Estilo NotebookLM)")

# 2. Inicializar cliente e historial en memoria
if "cliente" not in st.session_state:
    st.session_state.cliente = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])
if "documentos_guardados" not in st.session_state:
    st.session_state.documentos_guardados = {} # Guardará {nombre_archivo: texto}

# --- BARRA LATERAL: PANEL DE FUENTES ---
st.sidebar.title("📄 Panel de Fuentes")
st.sidebar.write("Sube tus documentos PDF de estudio.")

archivos_subidos = st.sidebar.file_uploader("Cargar PDFs", type=["pdf"], accept_multiple_files=True)

if archivos_subidos:
    for archivo in archivos_subidos:
        if archivo.name not in st.session_state.documentos_guardados:
            lector_pdf = PyPDF2.PdfReader(io.BytesIO(archivo.read()))
            texto_extraido = ""
            for pagina in lector_pdf.pages:
                texto_extraido += pagina.extract_text() + "\n"
            st.session_state.documentos_guardados[archivo.name] = texto_extraido

st.sidebar.divider()

# --- PANEL CENTRAL: CONFIGURACIÓN DE LA PRESENTACIÓN ---
st.write("### 🎛️ Configura tu Presentación")

if not st.session_state.documentos_guardados:
    st.warning("⚠️ Sube al menos un documento PDF en la barra lateral para comenzar.")
else:
    # Selector estilo NotebookLM: elegir qué fuentes usar
    st.write("#### 1. Selecciona las fuentes para esta presentación:")
    fuentes_seleccionadas = []
    
    col_sel1, col_sel2 = st.columns(2)
    for i, nombre_archivo in enumerate(st.session_state.documentos_guardados.keys()):
        col = col_sel1 if i % 2 == 0 else col_sel2
        if col.checkbox(nombre_archivo, value=True, key=f"src_{nombre_archivo}"):
            fuentes_seleccionadas.append(nombre_archivo)

    st.divider()

    # Parámetros de la presentación
    st.write("#### 2. Personaliza el diseño y enfoque")
    col_form1, col_form2 = st.columns(2)
    
    formato = col_form1.radio("Formato", ["Presentación detallada", "Diapositivas ejecutivas / resumen"])
    idioma = col_form2.selectbox("Seleccionar idioma", ["español", "inglés"])
    duracion = col_form2.selectbox("Duración aproximada", ["Corto (5 diapositivas)", "Estándar (10 diapositivas)", "Extenso (15 diapositivas)"])
    
    descripcion_prompt = st.text_area(
        "¿Qué enfoque específico quieres darle a la presentación?", 
        placeholder="Ej: Enfócate en los mecanismos fisiopatológicos y vías de señalización, estructurado de forma muy clara..."
    )

    if st.button("🚀 Generar Presentación PowerPoint (.pptx)", type="primary", use_container_width=True):
        if not fuentes_seleccionadas:
            st.error("Debes seleccionar al menos una fuente en la sección superior.")
        else:
            # Unir el texto de las fuentes seleccionadas
            texto_fuentes_unido = ""
            for nombre in fuentes_seleccionadas:
                texto_fuentes_unido += f"\n--- Fuente: {nombre} ---\n" + st.session_state.documentos_guardados[nombre]

            with st.spinner("🧠 Analizando fuentes y estructurando diapositivas con Gemini..."):
                prompt_ia = f"""
                Actúa como un diseñador instruccional y experto médico. 
                Utiliza ÚNICAMENTE la siguiente información de las fuentes seleccionadas: {texto_fuentes_unido[:10000]}
                
                Instrucciones de diseño y contenido:
                - Formato: {formato}
                - Idioma: {idioma}
                - Duración: {duracion}
                - Enfoque del usuario: {descripcion_prompt}
                
                Genera la estructura exacta de las diapositivas siguiendo estrictamente este formato de texto para que un script pueda leerlo:
                
                SLIDE: [Título de la diapositiva]
                - [Punto clave 1]
                - [Punto clave 2]
                - [Punto clave 3]
                
                (Repite la estructura SLIDE con sus viñetas para cada diapositiva según la duración solicitada. No inventes información que no esté en las fuentes).
                """

                try:
                    respuesta = st.session_state.cliente.models.generate_content(
                        model='gemini-3.8-flash',
                        contents=prompt_ia
                    )
                    
                    contenido_generado = respuesta.text
                    
                    # --- CONVERTIR TEXTO A POWERPOINT REAL (.pptx) ---
                    prs = Presentation()
                    
                    # Parsear el texto devuelto por la IA
                    bloques = contenido_generado.split("SLIDE:")
                    for bloque in bloques:
                        if not bloque.strip():
                            continue
                        
                        lineas = bloque.strip().split("\n")
                        titulo_slide = lineas[0].strip()
                        puntos = [l.strip("- ").strip() for l in lineas[1:] if l.strip().startswith("-")]
                        
                        # Crear la diapositiva en blanco o con título y objetos
                        slide_layout = prs.slide_layouts[1] # Layout de Título y Contenido
                        slide = prs.slides.add_slide(slide_layout)
                        
                        # Título
                        title_shape = slide.shapes.title
                        title_shape.text = titulo_slide
                        
                        # Cuerpo / Viñetas
                        body_shape = slide.placeholders[1]
                        tf = body_shape.text_frame
                        tf.text = "" # limpiar por defecto
                        
                        for i, punto in enumerate(puntos):
                            if i == 0:
                                tf.text = punto
                            else:
                                p = tf.add_paragraph()
                                p.text = punto
                                p.level = 0

                    # Guardar en memoria BytesIO para descargar
                    ppt_stream = io.BytesIO()
                    prs.save(ppt_stream)
                    ppt_stream.seek(0)

                    st.success("¡Presentación creada con éxito!")
                    
                    # Botón de descarga directa
                    st.download_button(
                        label="📥 Descargar tu archivo de PowerPoint (.pptx)",
                        data=ppt_stream,
                        file_name="Presentacion_Medica_Pro.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        type="primary"
                    )
                    
                    with st.expander("Ver vista previa del contenido estructurado"):
                        st.markdown(contenido_generado)

                except Exception as e:
                    st.error(f"Fallo al generar la presentación: {e}")